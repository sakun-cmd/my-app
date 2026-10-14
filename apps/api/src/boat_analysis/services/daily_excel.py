import logging
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from threading import Lock
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..config import get_settings
from ..database import engine
from ..models import Race, RaceEntry, RaceResult
from .official_results import NO_RESULT_MARKERS
from .scoring import ScoringInput, score_entries
from .two_boat_flow import analyze_races_two_boat_flow

logger = logging.getLogger(__name__)

_EXPORT_LOCK = Lock()
_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HIT_FILL = PatternFill("solid", fgColor="C6EFCE")
_DIFF_FILL = PatternFill("solid", fgColor="FFEB9C")


class DailyExcelNotReady(RuntimeError):
    pass


def daily_excel_path(race_date: date) -> Path:
    output_dir = get_settings().daily_excel_output_dir
    return output_dir / f"boat-race-daily-{race_date:%Y%m%d}.xlsx"


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "はい" if value else "いいえ"
    return value


def _latest_before(items: list[Any], cutoff: datetime | None) -> Any | None:
    if cutoff is None:
        return None
    eligible = [item for item in items if item.observed_at <= cutoff]
    return eligible[-1] if eligible else None


def _flow_columns(flow: dict[str, Any] | None, race: Race) -> dict[str, Any]:
    result = race.result
    winning = result.trifecta_combination if result else None
    tickets = flow.get("tickets", []) if flow else []
    ticket_by_combination = {
        "-".join(str(lane) for lane in ticket["combination"]): ticket for ticket in tickets
    }
    hit_ticket = ticket_by_combination.get(winning) if winning else None
    stake = len(tickets) * 100
    payout = result.trifecta_payout_yen if result else None
    # Cancelled / void races are refunded in full, so they must not count as losses.
    refunded = bool(
        result and (result.trifecta_refunded or result.winning_technique in NO_RESULT_MARKERS)
    )
    return_yen = stake if refunded else payout if hit_ticket and payout is not None else 0
    if not result or result.finalized_at is None:
        hit_status = "結果未確定"
    elif not tickets:
        hit_status = "買い目なし"
    elif refunded:
        hit_status = "返還"
    elif hit_ticket:
        hit_status = "的中"
    else:
        hit_status = "不的中"
    return {
        "判定状態": flow.get("status") if flow else "データなし",
        "注目艇": ",".join(str(lane) for lane in flow.get("likely_lanes", [])) if flow else "",
        "外艇": flow.get("outer_lane") if flow else None,
        "外の外艇": flow.get("outer_outer_lane") if flow else None,
        "イン逃げ評価": flow.get("escape_assessment") if flow else None,
        "注意喚起": (
            "イン逃げ注意（買い目作成なし）"
            if flow and flow.get("escape_assessment") == "unlikely"
            else None
        ),
        "最終判断": flow.get("decision") if flow else None,
        "買い目": " / ".join(ticket_by_combination),
        "買い目詳細": " / ".join(
            (
                f"{combination}"
                f"(オッズ={ticket.get('odds') or '-'},人気={ticket.get('popularity_rank') or '-'})"
            )
            for combination, ticket in ticket_by_combination.items()
        ),
        "買い目数": len(tickets),
        "投資額_各100円": stake,
        "合成オッズ": flow.get("composite_odds") if flow else None,
        "判定手順": " | ".join(
            f"{step['number']}:{step['result']} {step['detail']}" for step in flow.get("steps", [])
        )
        if flow
        else "",
        "3連単1着支持出典": flow.get("win_odds_source") if flow else None,
        "人気順位出典": flow.get("popularity_source") if flow else None,
        "公式3連単取得時刻": flow.get("win_odds_observed_at") if flow else None,
        "的中判定": hit_status,
        "的中買い目": winning if hit_ticket else None,
        "的中買い目オッズ": hit_ticket.get("odds") if hit_ticket else None,
        "的中買い目人気": hit_ticket.get("popularity_rank") if hit_ticket else None,
        "払戻額_100円": return_yen,
        "収支_各100円": return_yen - stake,
    }


def _combination_text(combination: tuple[int, ...] | list[int] | None) -> str | None:
    return "-".join(str(lane) for lane in combination) if combination else None


def _value_diffs(source: dict[int, Decimal | None], live: dict[int, Decimal | None]) -> list[str]:
    diffs = []
    for lane in sorted(set(source) | set(live)):
        before, after = source.get(lane), live.get(lane)
        if before == after:
            continue
        if before is None:
            diffs.append(f"{lane}:BT欠損")
        elif after is None:
            diffs.append(f"{lane}:実欠損")
        else:
            diffs.append(f"{lane}:{after - before:+.2f}")
    return diffs


def _max_abs_diff(
    source: dict[int, Decimal | None], live: dict[int, Decimal | None]
) -> Decimal | None:
    gaps = [
        abs(live[lane] - value)
        for lane, value in source.items()
        if value is not None and live.get(lane) is not None
    ]
    return max(gaps) if gaps else None


def _comparison_row(race: Race, flow: dict[str, Any] | None) -> dict[str, Any] | None:
    """バックテスト入力（元DBのみ）と実データ入力（公式値優先）の判定を1レース1行で比較する。"""
    baseline = flow.get("source_baseline") if flow else None
    if flow is None or baseline is None:
        return None
    source_flow = baseline["flow"]
    source_boats = baseline["boats"]
    live_metrics = {metric["lane"]: metric for metric in flow.get("metrics", [])}

    source_exhibition = {item.lane: item.exhibition_time for item in source_boats}
    live_exhibition = {lane: m.get("exhibition_time") for lane, m in live_metrics.items()}
    source_lap = {item.lane: item.lap_time for item in source_boats}
    live_lap = {lane: m.get("lap_time") for lane, m in live_metrics.items()}
    source_start_rank = {item.lane: item.course_avg_start_rank for item in source_boats}
    live_start_rank = {
        lane: m.get("course_avg_start_rank") for lane, m in live_metrics.items()
    }
    exhibition_diffs = _value_diffs(source_exhibition, live_exhibition)
    lap_diffs = _value_diffs(source_lap, live_lap)
    start_rank_diffs = _value_diffs(source_start_rank, live_start_rank)
    half_lap_used = any(
        observation.lap_source == "race.boatcast.jp"
        and observation.half_lap_time_seconds is not None
        for entry in race.entries
        for observation in entry.exhibitions
    )

    source_tickets = {tuple(row["combination"]) for row in source_flow.get("tickets", [])}
    live_tickets = {tuple(row["combination"]) for row in flow.get("tickets", [])}

    source_winning = baseline["winning_combination"]
    source_winning_text = _combination_text(source_winning)
    official_winning_text = race.result.trifecta_combination if race.result else None
    finish_comparable = source_winning_text is not None and official_winning_text is not None

    source_stake = len(source_tickets) * 100
    source_hit: bool | None = None
    source_return: Decimal | None = None
    if source_winning is not None:
        source_hit = bool(source_tickets) and source_winning in source_tickets
        odds = baseline["trifecta_odds"]
        if all(odds.get(ticket) for ticket in source_tickets):
            source_return = Decimal("100") * odds[source_winning] if source_hit else Decimal("0")

    live = _flow_columns(flow, race)
    live_hit = {"的中": True, "不的中": False, "買い目なし": False}.get(live["的中判定"])
    live_evaluable = live["的中判定"] not in {"結果未確定", "返還"}

    causes = []
    decision_differs = (
        source_flow.get("status") != flow.get("status") or source_tickets != live_tickets
    )
    if decision_differs:
        if exhibition_diffs:
            causes.append("展示タイム差")
        if lap_diffs:
            causes.append("周回タイム差")
        if start_rank_diffs:
            causes.append("コースST再計算差")
        if not exhibition_diffs and not lap_diffs and not start_rank_diffs:
            causes.append("判定差(入力差なし)")
    if finish_comparable and source_winning_text != official_winning_text:
        causes.append("着順データ差")
    if (
        source_hit
        and live_hit
        and source_return is not None
        and source_return != Decimal(live["払戻額_100円"])
    ):
        causes.append("オッズ・払戻差")

    return {
        "開催日": race.race_date,
        "場コード": race.venue_code,
        "場名": race.venue.name,
        "レース": race.race_number,
        "展示差_艇別(実-BT)": " ".join(exhibition_diffs),
        "展示差_最大秒": _max_abs_diff(source_exhibition, live_exhibition),
        "周回差_艇別(実-BT)": " ".join(lap_diffs),
        "周回差_最大秒": _max_abs_diff(source_lap, live_lap),
        "実周回種別": "半周" if half_lap_used else "一周",
        "BT判定状態": source_flow.get("status"),
        "実判定状態": flow.get("status"),
        "判定一致": source_flow.get("status") == flow.get("status"),
        "BT注目艇": ",".join(str(lane) for lane in source_flow.get("likely_lanes", [])),
        "実注目艇": ",".join(str(lane) for lane in flow.get("likely_lanes", [])),
        "BTイン逃げ評価": source_flow.get("escape_assessment"),
        "実イン逃げ評価": flow.get("escape_assessment"),
        "BT買い目数": len(source_tickets),
        "実買い目数": len(live_tickets),
        "買い目一致": source_tickets == live_tickets,
        "BTのみ買い目": " / ".join(
            _combination_text(ticket) or "" for ticket in sorted(source_tickets - live_tickets)
        ),
        "実のみ買い目": " / ".join(
            _combination_text(ticket) or "" for ticket in sorted(live_tickets - source_tickets)
        ),
        "BT着順(元DB)": source_winning_text,
        "公式3連単結果": official_winning_text,
        "着順一致": (source_winning_text == official_winning_text if finish_comparable else None),
        "BT的中": source_hit,
        "実的中判定": live["的中判定"],
        "的中一致": (
            source_hit == live_hit
            if source_hit is not None and live_hit is not None and live_evaluable
            else None
        ),
        "BT投資額": source_stake,
        "BT払戻額(元DBオッズ)": source_return,
        "BT収支": source_return - source_stake if source_return is not None else None,
        "実投資額": live["投資額_各100円"],
        "実払戻額(公式払戻)": live["払戻額_100円"],
        "実収支": live["収支_各100円"],
        "収支差(実-BT)": (
            live["収支_各100円"] - (source_return - source_stake)
            if source_return is not None and live_evaluable
            else None
        ),
        "相違要因": " / ".join(causes) if causes else "なし",
    }


def _comparison_summary(rows: list[dict[str, Any]]) -> list[tuple[str, Any]]:
    def ratio(key: str) -> str:
        values = [row[key] for row in rows if row[key] is not None]
        matched = sum(1 for value in values if value)
        return f"{matched}/{len(values)}" if values else "対象なし"

    def total(key: str) -> Any:
        values = [row[key] for row in rows if row[key] is not None]
        return sum(values, Decimal("0")) if values else None

    source_covered = [row for row in rows if row["BT払戻額(元DBオッズ)"] is not None]
    source_stake = sum(row["BT投資額"] for row in source_covered)
    source_return = sum((row["BT払戻額(元DBオッズ)"] for row in source_covered), Decimal("0"))
    live_evaluated = [
        row
        for row in rows
        if row["実的中判定"] not in {"結果未確定", "返還"} and row["実払戻額(公式払戻)"] is not None
    ]
    live_stake = sum(row["実投資額"] for row in live_evaluated)
    live_return = sum(row["実払戻額(公式払戻)"] for row in live_evaluated)
    cause_counts: dict[str, int] = {}
    for row in rows:
        for cause in row["相違要因"].split(" / "):
            cause_counts[cause] = cause_counts.get(cause, 0) + 1

    summary: list[tuple[str, Any]] = [
        ("比較対象レース", len(rows)),
        (
            "公式値で入力が変わったレース",
            sum(1 for row in rows if row["展示差_艇別(実-BT)"] or row["周回差_艇別(実-BT)"]),
        ),
        ("判定状態一致", ratio("判定一致")),
        ("買い目一致", ratio("買い目一致")),
        ("着順一致", ratio("着順一致")),
        ("的中一致", ratio("的中一致")),
        ("BT投資額(オッズ収録レース)", source_stake),
        ("BT払戻額", source_return),
        ("BT回収率", float(source_return / source_stake) if source_stake else None),
        ("実投資額(確定・返還除く)", live_stake),
        ("実払戻額", live_return),
        ("実回収率", live_return / live_stake if live_stake else None),
        ("収支差合計(実-BT)", total("収支差(実-BT)")),
    ]
    summary.extend((f"相違要因: {cause}", count) for cause, count in sorted(cause_counts.items()))
    return summary


def _race_row(race: Race, flow: dict[str, Any] | None) -> dict[str, Any]:
    weather = _latest_before(race.weather_observations, race.scheduled_at)
    result_by_lane = {entry.lane: entry for entry in race.result.entries} if race.result else {}
    metric_by_lane = {metric["lane"]: metric for metric in flow.get("metrics", [])} if flow else {}
    latest_exhibitions = {
        entry.lane: _latest_before(entry.exhibitions, race.scheduled_at) for entry in race.entries
    }
    scored = score_entries(
        [
            ScoringInput(
                lane=entry.lane,
                exhibition_time_seconds=(
                    latest_exhibitions[entry.lane].exhibition_time_seconds
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                exhibition_start_timing=(
                    latest_exhibitions[entry.lane].exhibition_start_timing
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                exhibition_course=(
                    latest_exhibitions[entry.lane].exhibition_course
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                half_lap_time_seconds=(
                    latest_exhibitions[entry.lane].half_lap_time_seconds
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                lap_time_seconds=(
                    latest_exhibitions[entry.lane].lap_time_seconds
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                cornering_time_seconds=(
                    latest_exhibitions[entry.lane].cornering_time_seconds
                    if latest_exhibitions[entry.lane]
                    else None
                ),
                straight_time_seconds=(
                    latest_exhibitions[entry.lane].straight_time_seconds
                    if latest_exhibitions[entry.lane]
                    else None
                ),
            )
            for entry in race.entries
        ],
        race.venue_code,
    )
    score_by_lane = {item.lane: item for item in scored}
    row: dict[str, Any] = {
        "開催日": race.race_date,
        "場コード": race.venue_code,
        "場名": race.venue.name,
        "レース": race.race_number,
        "締切予定": race.scheduled_at,
        "グレード": race.grade,
        "距離m": race.distance_m,
        "結果確定時刻": race.result.finalized_at if race.result else None,
        "結果出典": race.result.source if race.result else None,
        "決まり手": race.result.winning_technique if race.result else None,
        "3連単結果": race.result.trifecta_combination if race.result else None,
        "3連単払戻円": race.result.trifecta_payout_yen if race.result else None,
        "3連単人気": race.result.trifecta_popularity_rank if race.result else None,
        "3連単返還": race.result.trifecta_refunded if race.result else None,
        "天候": weather.weather if weather else None,
        "気温C": weather.air_temperature_c if weather else None,
        "水温C": weather.water_temperature_c if weather else None,
        "風向": weather.wind_direction if weather else None,
        "風速m_s": weather.wind_speed_mps if weather else None,
        "波高cm": weather.wave_height_cm if weather else None,
    }
    row.update(_flow_columns(flow, race))

    entries = {entry.lane: entry for entry in race.entries}
    for lane in range(1, 7):
        entry = entries.get(lane)
        prefix = f"{lane}号艇_"
        exhibition = latest_exhibitions.get(lane)
        win_odds = _latest_before(entry.win_odds_observations, race.scheduled_at) if entry else None
        metric = metric_by_lane.get(lane, {})
        score = score_by_lane.get(lane)
        race_result = result_by_lane.get(lane)
        replacements = (
            " / ".join(f"{part.part_name}x{part.quantity}" for part in entry.part_replacements)
            if entry
            else ""
        )
        row.update(
            {
                f"{prefix}登録番号": entry.racer_registration_number if entry else None,
                f"{prefix}選手名": entry.racer.name if entry and entry.racer else None,
                f"{prefix}級別": entry.class_name if entry else None,
                f"{prefix}モーター": entry.motor_number if entry else None,
                f"{prefix}ボート": entry.boat_number if entry else None,
                f"{prefix}体重kg": entry.weight_kg if entry else None,
                f"{prefix}3連単1着支持換算値": (
                    win_odds.win_odds if win_odds else metric.get("win_odds")
                ),
                f"{prefix}人気順位": metric.get("popularity_rank"),
                f"{prefix}展示進入": exhibition.exhibition_course if exhibition else None,
                f"{prefix}展示ST": exhibition.exhibition_start_timing if exhibition else None,
                f"{prefix}展示タイム": exhibition.exhibition_time_seconds if exhibition else None,
                f"{prefix}展示補正後": metric.get("exhibition_time_corrected"),
                f"{prefix}展示順位": metric.get("exhibition_rank"),
                f"{prefix}チルト": exhibition.tilt if exhibition else None,
                f"{prefix}半周": exhibition.half_lap_time_seconds if exhibition else None,
                f"{prefix}一周": exhibition.lap_time_seconds if exhibition else None,
                f"{prefix}周回比較値": metric.get("lap_time"),
                f"{prefix}周回補正後": metric.get("lap_time_corrected"),
                f"{prefix}周回順位": metric.get("lap_rank"),
                f"{prefix}まわり足": exhibition.cornering_time_seconds if exhibition else None,
                f"{prefix}直線": exhibition.straight_time_seconds if exhibition else None,
                f"{prefix}展示出典": exhibition.source if exhibition else None,
                f"{prefix}周回出典": exhibition.lap_source if exhibition else None,
                f"{prefix}AP": metric.get("ap"),
                f"{prefix}参考スコア": score.score if score else None,
                f"{prefix}参考順位": score.rank if score else None,
                f"{prefix}信頼度": score.confidence if score else None,
                f"{prefix}警告": " / ".join(score.warnings) if score else None,
                f"{prefix}コース平均ST順位": metric.get("course_avg_start_rank"),
                f"{prefix}コース平均ST": metric.get("course_avg_start_timing"),
                f"{prefix}全国勝率": metric.get("nationwide_win_rate"),
                f"{prefix}コース3連率": metric.get("course_top3_rate"),
                f"{prefix}コース進入数": metric.get("course_starts"),
                f"{prefix}コース成績基準日": metric.get("course_stats_snapshot_date"),
                f"{prefix}コース成績出典": metric.get("course_stats_source"),
                f"{prefix}部品交換": replacements,
                f"{prefix}着順": race_result.finish_position if race_result else None,
                f"{prefix}本番進入": race_result.actual_course if race_result else None,
                f"{prefix}本番ST": race_result.actual_start_timing if race_result else None,
                f"{prefix}結果状態": race_result.status if race_result else None,
            }
        )
    return row


def _append_comparison_sheets(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("バックテスト比較")
    headers = list(rows[0]) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(row.get(header)) for header in headers])
    sheet.freeze_panes = "E2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cause_column = headers.index("相違要因") + 1 if "相違要因" in headers else None
    if cause_column:
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=cause_column).value != "なし":
                for cell in sheet[row_number]:
                    cell.fill = _DIFF_FILL
    for index, header in enumerate(headers, start=1):
        width = 14
        if "艇別" in header or header.endswith("のみ買い目"):
            width = 30
        elif header == "相違要因":
            width = 36
        sheet.column_dimensions[get_column_letter(index)].width = width

    summary = workbook.create_sheet("比較集計")
    summary.append(["項目", "値"])
    for label, value in _comparison_summary(rows) if rows else [("比較対象レース", 0)]:
        summary.append([label, _excel_value(value)])
    for cell in summary[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 18


def build_daily_workbook(
    rows: list[dict[str, Any]],
    race_date: date,
    generated_at: datetime,
    comparison_rows: list[dict[str, Any]] | None = None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "レース結果"
    headers = list(rows[0]) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(row.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hit_column = headers.index("的中判定") + 1 if "的中判定" in headers else None
    if hit_column:
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=hit_column).value == "的中":
                for cell in sheet[row_number]:
                    cell.fill = _HIT_FILL
    for index, header in enumerate(headers, start=1):
        width = 14
        if header in {"買い目", "買い目詳細", "判定手順"}:
            width = 42
        elif header.endswith("選手名") or header.endswith("部品交換"):
            width = 18
        sheet.column_dimensions[get_column_letter(index)].width = width

    info = workbook.create_sheet("説明")
    info.append(["項目", "内容"])
    info.append(["対象日", race_date.isoformat()])
    info.append(["生成時刻", generated_at.isoformat()])
    info.append(["単位", "買い目は各100円。払戻額は公式3連単の100円払戻。"])
    info.append(
        [
            "買い目",
            "2艇判定フローを締切前に保存された展示・周回・公式3連単オッズから日次再計算。",
        ]
    )
    info.append(
        [
            "バックテスト比較",
            "BT列はバックテストと同じ元DBの展示(brd_c3)・一周(brd_c4)・着順(brd_k3)・"
            "3連単オッズ(odds_view)だけで判定。実列は公式展示・Boatcast周回・公式結果・"
            "公式払戻を優先。差は実-BT。相違があるレースは黄色で表示。",
        ]
    )
    info.append(
        [
            "注意",
            "統計的な判断支援用です。的中・利益を保証せず、自動投票は行いません。",
        ]
    )
    for cell in info[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 90
    if comparison_rows is not None:
        _append_comparison_sheets(workbook, comparison_rows)
    return workbook


def generate_daily_excel(
    race_date: date,
    *,
    require_complete: bool = True,
    force: bool = False,
) -> Path:
    path = daily_excel_path(race_date)
    if path.exists() and not force:
        return path
    if not _EXPORT_LOCK.acquire(blocking=False):
        raise RuntimeError("daily Excel export is already running")
    try:
        with Session(engine) as session:
            races = session.scalars(
                select(Race)
                .where(Race.race_date == race_date)
                .options(
                    selectinload(Race.venue),
                    selectinload(Race.entries).selectinload(RaceEntry.racer),
                    selectinload(Race.entries).selectinload(RaceEntry.exhibitions),
                    selectinload(Race.entries).selectinload(RaceEntry.part_replacements),
                    selectinload(Race.entries).selectinload(RaceEntry.win_odds_observations),
                    selectinload(Race.weather_observations),
                    selectinload(Race.result).selectinload(RaceResult.entries),
                )
                .order_by(Race.scheduled_at, Race.venue_code, Race.race_number)
            ).all()
            if not races:
                raise DailyExcelNotReady(f"{race_date.isoformat()}のレースがありません")
            incomplete = [
                race
                for race in races
                if race.result is None
                or race.result.finalized_at is None
                or (
                    race.result.winning_technique not in NO_RESULT_MARKERS
                    and race.result.trifecta_payout_yen is None
                )
            ]
            if require_complete and incomplete:
                raise DailyExcelNotReady(
                    f"{race_date.isoformat()}は未確定または払戻未取得が{len(incomplete)}レースあります"
                )
            flows = analyze_races_two_boat_flow(races, session, include_source_baseline=True)
            rows = [_race_row(race, flows.get(race.id)) for race in races]
            comparison_rows = [
                row
                for race in races
                if (row := _comparison_row(race, flows.get(race.id))) is not None
            ]

        generated_at = datetime.now(ZoneInfo(get_settings().app_timezone))
        workbook = build_daily_workbook(rows, race_date, generated_at, comparison_rows)
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(".tmp.xlsx")
        workbook.save(temporary)
        temporary.replace(path)
        logger.info("Daily Excel exported date=%s races=%s", race_date, len(rows))
        return path
    finally:
        _EXPORT_LOCK.release()


def generate_due_daily_excel(now: datetime | None = None) -> Path | None:
    settings = get_settings()
    if not settings.daily_excel_enabled:
        return None
    timezone = ZoneInfo(settings.app_timezone)
    current = now or datetime.now(timezone)
    current = (
        current.replace(tzinfo=timezone) if current.tzinfo is None else current.astimezone(timezone)
    )
    scheduled = datetime.combine(
        current.date(),
        time(settings.daily_excel_hour_jst, settings.daily_excel_minute_jst),
        timezone,
    )
    if current < scheduled:
        return None
    target_date = current.date() - timedelta(days=1)
    path = daily_excel_path(target_date)
    if path.exists():
        return path
    return generate_daily_excel(target_date)
